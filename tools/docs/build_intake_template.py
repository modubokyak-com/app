# 연구원용 조합 추가 입력 양식(docs/interaction-rule-intake-template.xlsx) 생성기.
# 입력: 약학검수시트.xlsx에서 뽑은 sheet.json (규칙 70·성분 149). 사용: python tools/docs/build_intake_template.py <sheet.json 폴더> <출력 xlsx>
import json, sys
from openpyxl import Workbook
from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
from openpyxl.worksheet.datavalidation import DataValidation
from openpyxl.utils import get_column_letter
from openpyxl.comments import Comment

SP = sys.argv[1]; OUT = sys.argv[2]
d = json.load(open(f"{SP}/sheet.json", encoding="utf-8"))
F = "맑은 고딕"
def font(**k): return Font(name=F, size=k.pop("size", 11), **k)
HEAD = PatternFill("solid", fgColor="1F3A5F"); HEADF = Font(name=F, size=11, bold=True, color="FFFFFF")
INPUT = PatternFill("solid", fgColor="FFF8C6")   # 연노랑 = 적는 칸
EX = PatternFill("solid", fgColor="EEF2F7")      # 회색 = 예시
thin = Side(style="thin", color="C8CED6"); BORDER = Border(left=thin, right=thin, top=thin, bottom=thin)
WRAP = Alignment(wrap_text=True, vertical="top")

wb = Workbook()

# ───────── 0. 읽어주세요
ws = wb.active; ws.title = "읽어주세요"
ws.column_dimensions["A"].width = 110
lines = [
 ("모두의 복약 — 약 × 영양제 조합 추가 입력 양식 (v1, 2026-10-12)", "title"),
 ("", None),
 ("왜 이 양식인가", "h"),
 ("앱의 「1분 복용 점검」은 데이터베이스에 들어 있는 조합 규칙만 보여 줍니다. 앱이 문장을 지어내지 않습니다.", None),
 ("그래서 연구원님이 적어 주신 「무슨 일이」「어떻게 하면」 문장이 그대로 사용자 화면에 나갑니다. 화면용 문장을 가장 신경 써 주세요.", None),
 ("", None),
 ("어디에 적나 — 노란 칸만 적으시면 됩니다", "h"),
 ("시트 ①「새 조합」 : 새로 넣을 약 × 영양제(또는 영양제 × 영양제) 조합. 한 줄 = 조합 하나. 회색 줄 3개는 예시이니 지우지 말고 그 아래부터 적어 주세요.", None),
 ("시트 ②「새 성분」 : 조합에 쓴 성분·계열 이름이 시트 ④ 목록에 없을 때만 적습니다. 시트 ①의 「목록 확인」 칸이 자동으로 알려 줍니다.", None),
 ("시트 ③「이미 있는 조합」 : 지금 데이터베이스에 있는 70건. 중복으로 적지 않도록 먼저 훑어봐 주세요. 고칠 것이 있으면 맨 오른쪽 「의견」 칸에 적어 주세요.", None),
 ("시트 ④「이미 있는 성분·계열」 : 이름 표기를 이 목록과 똑같이 맞춰 주세요 (예: 오메가-3, 비타민 D, 코엔자임 Q10).", None),
 ("시트 ⑤「용어」 : 등급·근거 수준·조건의 뜻과 고를 수 있는 값. 시트 ①의 드롭다운이 이 시트를 읽습니다.", None),
 ("", None),
 ("꼭 지켜 주실 것 5가지", "h"),
 ("1. 한 줄에 조합 하나. 「A × B 또는 C」처럼 여러 조합을 한 줄에 적지 말고 줄을 나눠 주세요. 단, 「PPI × (마그네슘 또는 비타민 B12)」처럼 한쪽이 '둘 중 하나'인 규칙은 B칸에 「마그네슘 / 비타민 B12」로 적고 비고에 '둘 중 하나'라고 써 주세요.", None),
 ("2. 성분·계열 이름은 시트 ④ 표기 그대로. 계열(스타틴, 항혈소판제 …)로 적어도 되고 개별 성분(아토르바스타틴)으로 적어도 됩니다. 계열로 적으면 그 계열의 모든 약에 걸립니다.", None),
 ("3. 등급·근거 수준은 드롭다운에서만 고릅니다. 뜻은 시트 ⑤에 있습니다. 근거가 얇은데 등급이 높으면 비고에 이유를 적어 주세요.", None),
 ("4. 화면용 문장(무슨 일이·어떻게 하면)은 70대 어르신이 읽는 글입니다. 각 2문장 이내, 전문용어(INR·CYP3A4 등)는 풀어 쓰고, 끝은 '~해 주세요' '~하세요'. 「담당 선생님과 상의」처럼 다음 행동이 보이게 적어 주세요.", None),
 ("5. 출처는 반드시. 책이면 「책이름 쪽수」, 논문이면 「저자 연도 저널」, 식약처 자료면 「고시번호」. 출처 없는 줄은 데이터베이스에 넣지 않고 보류합니다.", None),
 ("", None),
 ("적지 않아도 되는 것", "h"),
 ("영문 코드, 데이터베이스 ID, 규칙 코드는 개발자가 붙입니다. 「시간 간격」「수술 전 중단」은 해당할 때만 적습니다.", None),
 ("", None),
 ("보내 주신 뒤 절차", "h"),
 ("개발자가 그대로 데이터베이스에 넣되 '검수 중' 상태로 두고, 약사 승인 뒤에만 앱에 나갑니다(데이터베이스 제약으로 승인 전 게시가 막혀 있습니다). 넣은 뒤 어떤 입력에 걸렸는지 결과 화면 캡처를 돌려 드립니다.", None),
]
for i,(t,k) in enumerate(lines, 1):
    c = ws.cell(row=i, column=1, value=t); c.alignment = WRAP
    if k=="title": c.font = font(size=16, bold=True, color="1F3A5F")
    elif k=="h": c.font = font(size=12, bold=True, color="1F3A5F")
    else: c.font = font()
ws.cell(row=len(lines)+2, column=1, value="칸 색 안내:   노란 칸 = 연구원님이 적는 곳      회색 줄 = 예시(지우지 마세요)      흰 칸 = 자동 계산·참고").font = font(italic=True, color="555555")

# ───────── 5. 용어 (먼저 만들어 두어야 드롭다운 참조 가능)
wt = wb.create_sheet("⑤ 용어")
grades = [
 ("금기", "contraindicated", "함께 쓰면 안 됨. 둘 중 하나를 빼야 함.", "아미오다론 × 요오드·켈프"),
 ("주의", "caution", "같이 쓰면 해로울 수 있어 의사·약사 상의가 필요함.", "와파린 × 크랜베리"),
 ("시간 조정", "timing", "같이 써도 되지만 복용 시간을 띄워야 함. 「시간 간격」 칸을 꼭 채움.", "레보티록신 × 칼슘 (4시간)"),
 ("확인 권장", "monitor", "함께 쓸 수 있으나 검사 수치·증상을 지켜봐야 함.", "스타틴 × 홍국"),
 ("알아두기", "info", "위험은 낮고 알아 두면 좋은 정보.", "비타민 C × 철분 (흡수 도움)"),
]
evid = [
 ("확립", "established", "임상 연구·공식 자료로 확인됨", "연구로 확인됨"),
 ("이론적", "theoretical", "기전상 예상되나 사람에서 확인되지 않음", "가능성 있음"),
 ("제한적", "limited", "사례 보고 몇 건 수준", "가능성 있음"),
 ("엇갈림", "conflicting", "연구 결과가 서로 다름", "의견이 갈림"),
]
conds = [
 ("임신", "앱 버튼 「임신·수유 중」이 이 조건으로 전달됨"),
 ("만성콩팥병(3~4기·중증)", "앱 버튼 「신장질환」은 아직 연결 안 됨 — 신장 조건 규칙을 적어 주시면 연결함"),
 ("혈액투석", ""),
 ("흡연", ""),
 ("수술 예정", "「수술 전 중단」 일수와 함께 쓰는 조건"),
 ("간질환", "데이터베이스에 아직 없음 — 적어 주시면 새로 만듦"),
 ("60대 이상", "데이터베이스에 아직 없음 — 고령 조건 규칙이 있으면 적어 주시면 새로 만듦"),
]
chips = ["오메가3","비타민D","마그네슘","유산균","종합비타민","철분","루테인","밀크씨슬","콜라겐","아연","홍삼","단백질보충제",
         "갑상선약","혈압약","고지혈증약","위장약","통증·소염제","알레르기약","피임약","항우울제","여드름약"]
chip_note = {"유산균":"= 프로바이오틱스","홍삼":"= 인삼","단백질보충제":"= 콩","갑상선약":"= 레보티록신","혈압약":"= 칼슘통로차단제·RAS차단제·칼륨보존이뇨제",
             "고지혈증약":"= 스타틴","위장약":"= PPI","항우울제":"= SSRI·SNRI","통증·소염제":"연결된 규칙 없음 — NSAID·아세트아미노펜 조합이 필요","알레르기약":"연결된 규칙 없음 — 항히스타민제 조합이 필요","피임약":"연결된 규칙 없음","여드름약":"연결된 규칙 없음 — 이소트레티노인 등","종합비타민":"연결된 규칙 없음","콜라겐":"연결된 규칙 없음"}
def table(ws, r0, title, headers, rows):
    ws.cell(row=r0, column=1, value=title).font = font(size=12, bold=True, color="1F3A5F")
    for j,h in enumerate(headers,1):
        c=ws.cell(row=r0+1, column=j, value=h); c.font=HEADF; c.fill=HEAD; c.border=BORDER; c.alignment=WRAP
    for i,row in enumerate(rows):
        for j,v in enumerate(row,1):
            c=ws.cell(row=r0+2+i, column=j, value=v); c.font=font(); c.border=BORDER; c.alignment=WRAP
    return r0+2+len(rows)
r = table(wt, 1, "A. 등급 (심각도) — 시트① 「등급」 드롭다운", ["등급","DB 값(개발자용)","뜻","예"], grades)
G0, G1 = 3, 2+len(grades)
r = table(wt, r+2, "B. 근거 수준 — 시트① 「근거 수준」 드롭다운", ["근거 수준","DB 값(개발자용)","뜻","화면에 보이는 표기"], evid)
E0, E1 = r-len(evid), r-1
r = table(wt, r+2, "C. 조건 (해당할 때만) — 시트① 「조건」 드롭다운", ["조건 이름","비고"], conds)
C0, C1 = r-len(conds), r-1
r = table(wt, r+2, "D. 앱 화면의 버튼(칩) — 사용자가 누르는 이름과 데이터베이스 성분·계열의 대응", ["앱 버튼","데이터베이스에서 풀리는 이름 / 상태"], [(c, chip_note.get(c,"같은 이름의 성분 있음")) for c in chips])
kinds = [("의약품","개별 약 성분 (와파린, 메트포르민)"),("약효군","약 계열 (스타틴, PPI, 항혈소판제)"),("생약·식물","홍국, 은행잎, 인삼 …"),("비타민",""),("미네랄",""),("프로바이오틱스",""),("기타","오메가-3, 코엔자임 Q10, 아미노산 …")]
r = table(wt, r+2, "E. 성분 종류 — 시트② 「종류」 드롭다운", ["종류","뜻"], kinds)
K0, K1 = r-len(kinds), r-1
for col,w in zip("ABCD",[26,22,70,34]): wt.column_dimensions[col].width = w

# ───────── 4. 이미 있는 성분·계열
w4 = wb.create_sheet("④ 이미 있는 성분·계열")
subs = d["subs"]
extra = [("고용량 항산화제","기타"),("디곡신","의약품"),("레보도파","의약품"),("베타카로틴","비타민"),("비타민 B12","비타민"),("사카로마이세스 보울라디","프로바이오틱스"),
         ("스타프루트","생약·식물"),("엽산","비타민"),("인","미네랄"),("콜린에스터라제 억제제","약효군"),("피나스테리드","의약품"),("항레트로바이러스제","약효군"),
         ("항생제","약효군"),("항콜린성 약물","약효군")]
seen = {s[0] for s in subs}
allsubs = sorted(set(tuple(s) for s in subs) | {e for e in extra if e[0] not in seen}, key=lambda x:(x[1], x[0]))
w4.cell(row=1, column=1, value="데이터베이스에 이미 있는 성분·계열 이름 (약학검수시트 v1 기준 %d개). 시트①에 적을 때 이 표기를 그대로 쓰세요." % len(allsubs)).font=font(bold=True, color="1F3A5F")
for j,h in enumerate(["이름 (이 표기 그대로)","종류"],1):
    c=w4.cell(row=2,column=j,value=h); c.font=HEADF; c.fill=HEAD; c.border=BORDER
for i,(n,k) in enumerate(allsubs):
    for j,v in enumerate((n,k),1):
        c=w4.cell(row=3+i,column=j,value=v); c.font=font(); c.border=BORDER
w4.column_dimensions["A"].width=34; w4.column_dimensions["B"].width=16; w4.freeze_panes="A3"
w4.auto_filter.ref = f"A2:B{2+len(allsubs)}"
SUBRANGE = f"'④ 이미 있는 성분·계열'!$A$3:$A${2+len(allsubs)}"

# ───────── 3. 이미 있는 조합
w3 = wb.create_sheet("③ 이미 있는 조합")
w3.cell(row=1,column=1,value="데이터베이스에 이미 있는 조합 70건 (2026-10 기준). 같은 조합은 시트①에 다시 적지 말고, 고칠 점이 있으면 「의견」에 적어 주세요.").font=font(bold=True,color="1F3A5F")
hdr=["번호","등급","근거 수준","조합","무슨 일이 (화면)","어떻게 하면 (화면)","시간 간격","출처","의견 (고칠 점이 있으면)"]
for j,h in enumerate(hdr,1):
    c=w3.cell(row=2,column=j,value=h); c.font=HEADF; c.fill=HEAD; c.border=BORDER; c.alignment=WRAP
for i,rw in enumerate(d["rules"]):
    vals=[i+1, rw[1], rw[2], rw[3], rw[5], rw[6], rw[7], rw[8], None]
    for j,v in enumerate(vals,1):
        c=w3.cell(row=3+i,column=j,value=v); c.font=font(); c.border=BORDER; c.alignment=WRAP
        if j==9: c.fill=INPUT
for col,w in zip("ABCDEFGHI",[6,10,10,30,50,50,10,16,36]): w3.column_dimensions[col].width=w
w3.freeze_panes="A3"; w3.auto_filter.ref=f"A2:I{2+len(d['rules'])}"

# ───────── 1. 새 조합
w1 = wb.create_sheet("① 새 조합", 1)
headers = [
 ("번호", 6, "순서대로 1, 2, 3 …"),
 ("A쪽: 약 (성분 또는 계열)", 24, "시트④ 표기 그대로. 계열(스타틴)도 가능"),
 ("B쪽: 영양제·건기식 (또는 두 번째 약)", 26, "시트④ 표기 그대로. '둘 중 하나'면 「마그네슘 / 비타민 B12」처럼 /로 구분"),
 ("조건 (해당할 때만)", 18, "임신·수술 예정 등. 드롭다운. 조건이 있을 때만 걸리는 규칙에 사용"),
 ("등급", 11, "드롭다운. 뜻은 시트⑤ A"),
 ("근거 수준", 11, "드롭다운. 뜻은 시트⑤ B"),
 ("요약 한 줄 (화면 제목)", 36, "결과 목록에 보이는 한 줄. 40자 이내"),
 ("무슨 일이 (화면)", 50, "어르신 눈높이 2문장 이내. 왜 조심해야 하는지"),
 ("어떻게 하면 (화면)", 50, "다음 행동. '~해 주세요'로 끝맺기"),
 ("시간 간격(시간)", 10, "등급이 「시간 조정」일 때 숫자만. 예: 4"),
 ("수술 전 중단(일)", 10, "수술 예정 조건일 때 숫자만. 예: 7"),
 ("출처", 24, "책 쪽수 / 논문 / 고시번호. 필수"),
 ("기전 요약 (내부용, 선택)", 40, "화면에 안 나감. 약사 검수용 메모"),
 ("비고", 30, "'둘 중 하나' 표시, 근거가 얇은 이유 등"),
 ("목록 확인 (자동)", 22, "자동 계산 — 손대지 마세요"),
 ("작성자 / 날짜", 16, ""),
]
w1.cell(row=1,column=1,value="새로 넣을 조합 — 노란 칸만 적어 주세요. 회색 3줄은 예시입니다(지우지 마세요). 「목록 확인」이 '새 성분'이라고 뜨면 시트②에 그 성분을 적어 주세요.").font=font(bold=True,color="1F3A5F")
for j,(h,w,tip) in enumerate(headers,1):
    c=w1.cell(row=2,column=j,value=h); c.font=HEADF; c.fill=HEAD; c.border=BORDER; c.alignment=WRAP
    w1.column_dimensions[get_column_letter(j)].width=w
    if tip: c.comment=Comment(tip, "모두의 복약")
w1.row_dimensions[2].height=46
examples = [
 ["예시1","와파린","크랜베리","","주의","제한적","와파린을 드시면 크랜베리는 피해 주세요",
  "크랜베리를 함께 드신 뒤 피가 잘 멎지 않는 쪽으로 수치가 올라갔다는 보고가 있습니다. 출혈 위험이 커질 수 있습니다.",
  "요로감염 예방으로 크랜베리를 생각하신다면, 와파린 복용 사실을 알리고 다른 방법이 있는지 담당 선생님과 상의해 주세요.",
  None,None,"임상건강기능식품학 부록 §4 p.358","크랜베리 플라보노이드가 CYP2C9 억제 → 와파린 농도↑, INR↑ 사례보고","이미 DB에 있는 규칙을 예시로 옮긴 것",None,"예시"],
 ["예시2","레보티록신","칼슘 / 철분","","시간 조정","확립","갑상선약과 칼슘·철분은 4시간 띄워 드세요",
  "칼슘이나 철분이 갑상선약을 장에서 붙잡아 흡수를 줄입니다. 효과가 약해질 수 있습니다.",
  "갑상선약은 아침 공복에 먼저 드시고, 칼슘·철분은 4시간 뒤에 드세요.",
  4,None,"임상건강기능식품학 부록 §4 p.372","다가양이온이 레보티록신과 불용성 복합체 형성","B쪽 '둘 중 하나' — 칼슘만 골라도, 철분만 골라도 걸림",None,"예시"],
 ["예시3","항혈소판제","은행잎","수술 예정","주의","확립","수술 전에는 은행잎을 미리 끊어 주세요",
  "은행잎은 피를 묽게 하는 쪽으로 작용합니다. 피를 묽게 하는 약과 겹치면 수술 중 출혈이 늘 수 있습니다.",
  "수술 날짜가 잡혀 있다면 적어도 1주일 전에는 은행잎 제품을 멈추고, 드시고 있다는 사실을 수술팀에 알려 주세요.",
  None,7,"임상건강기능식품학 부록 §4 p.377","은행잎 테르펜락톤의 PAF 길항 → 혈소판 응집 억제 가산","조건이 있는 규칙의 예 — 수술 예정일 때만 이 문구가 나감",None,"예시"],
]
N_INPUT = 40
for i,ex in enumerate(examples):
    rr=3+i
    for j,v in enumerate(ex,1):
        c=w1.cell(row=rr,column=j,value=v); c.font=font(color="555555"); c.fill=EX; c.border=BORDER; c.alignment=WRAP
    w1.row_dimensions[rr].height=96
first = 3+len(examples); last = first+N_INPUT-1
for rr in range(first,last+1):
    for j in range(1,len(headers)+1):
        c=w1.cell(row=rr,column=j); c.border=BORDER; c.alignment=WRAP; c.font=font()
        if j not in (1,15): c.fill=INPUT
    w1.cell(row=rr,column=1,value=rr-first+1)
# 목록 확인 공식: A쪽과 B쪽(/ 앞 첫 항목)이 시트④에 있는지
for rr in range(3,last+1):
    a=f"B{rr}"; b=f"C{rr}"
    bfirst=f'TRIM(IF(ISNUMBER(FIND("/",{b})),LEFT({b},FIND("/",{b})-1),{b}))'
    f=(f'=IF(AND({a}="",{b}=""),"",'
       f'IF(AND(OR({a}="",ISNUMBER(MATCH(TRIM({a}),{SUBRANGE},0))),OR({b}="",ISNUMBER(MATCH({bfirst},{SUBRANGE},0)))),"둘 다 목록에 있음",'
       f'"새 성분 → 시트②에 적어 주세요"))')
    c=w1.cell(row=rr,column=15,value=f); c.font=font(italic=True, color="1F3A5F")
def dv(ws, formula, cells):
    v=DataValidation(type="list", formula1=formula, allow_blank=True, showErrorMessage=True,
                     errorTitle="목록에서 골라 주세요", error="시트⑤ 용어의 값만 쓸 수 있습니다.")
    ws.add_data_validation(v); v.add(cells)
dv(w1, f"'⑤ 용어'!$A${G0}:$A${G1}", f"E{first}:E{last}")
dv(w1, f"'⑤ 용어'!$A${E0}:$A${E1}", f"F{first}:F{last}")
dv(w1, f"'⑤ 용어'!$A${C0}:$A${C1}", f"D{first}:D{last}")
w1.freeze_panes="D3"

# ───────── 2. 새 성분
w2 = wb.create_sheet("② 새 성분", 2)
w2.cell(row=1,column=1,value="시트④ 목록에 없는 성분·계열만 적어 주세요. 영문명·다른 표기는 사용자가 직접 글자로 입력할 때 찾아 주는 데 쓰입니다.").font=font(bold=True,color="1F3A5F")
h2=[("한글 이름 (대표 표기)",26,"시트①에 적은 것과 똑같이"),("영문 이름",24,"일반명. 예: cranberry, duloxetine"),
    ("다른 표기 (쉼표로 구분)",36,"염 표기·흔한 상품명·오타. 예: 둘록세틴염산염, 심발타"),
    ("종류",14,"드롭다운. 시트⑤ E"),("속하는 계열 (있으면)",24,"예: 아토르바스타틴 → 스타틴. 계열도 시트④ 표기로"),
    ("앱 버튼으로 바로 보이게 할까요?",22,"예/아니오. 예면 어떤 이름의 버튼이면 좋을지"),("출처·비고",36,"")]
for j,(h,w,tip) in enumerate(h2,1):
    c=w2.cell(row=2,column=j,value=h); c.font=HEADF; c.fill=HEAD; c.border=BORDER; c.alignment=WRAP
    w2.column_dimensions[get_column_letter(j)].width=w
    if tip: c.comment=Comment(tip,"모두의 복약")
w2.row_dimensions[2].height=40
ex2=["둘록세틴 (예시 — 이미 있음)","duloxetine","둘록세틴염산염, 심발타","의약품","SNRI","아니오 — 「항우울제」 버튼이 SNRI로 풀림","예시 줄"]
for j,v in enumerate(ex2,1):
    c=w2.cell(row=3,column=j,value=v); c.font=font(color="555555"); c.fill=EX; c.border=BORDER; c.alignment=WRAP
for rr in range(4,4+25):
    for j in range(1,len(h2)+1):
        c=w2.cell(row=rr,column=j); c.fill=INPUT; c.border=BORDER; c.alignment=WRAP; c.font=font()
dv(w2, f"'⑤ 용어'!$A${K0}:$A${K1}", "D4:D28")
w2.freeze_panes="A3"

order=["읽어주세요","① 새 조합","② 새 성분","③ 이미 있는 조합","④ 이미 있는 성분·계열","⑤ 용어"]
wb._sheets=[wb[n] for n in order]
from openpyxl.workbook.properties import CalcProperties
wb.calculation=CalcProperties(fullCalcOnLoad=True)
wb.save(OUT); print("saved", OUT)
